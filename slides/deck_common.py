"""Shared look and data access for the weekly supervisor decks (weeks 2-6).

Same colours, fonts and helpers as build_week01.py. Every result number shown on a slide is read
from reports/tables/*.csv (written by `python -m medical_tokenizer.report`) and formatted here; no
result is typed by hand in the builder scripts.
"""
import csv
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "reports" / "tables"
FIGS = ROOT / "reports" / "figures"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
GREY = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1E, 0x29, 0x3B)
RED = RGBColor(0xB9, 0x1C, 0x1C)
AMBER = RGBColor(0xB4, 0x53, 0x09)
MINT = RGBColor(0x5E, 0xEA, 0xD4)
PALE = RGBColor(0xCB, 0xD5, 0xE1)
FONT = "Calibri"

# the six-week plan, as agreed with the supervisor (slides/WEEKLY_PLAN.md)
PLAN = [
    ("Week 1", "Oct 8", "Foundation and protocol"),
    ("Week 2", "Oct 15", "Image tokenization (MedMNIST)"),
    ("Week 3", "Oct 22", "Time series (PhysioNet 2012)"),
    ("Week 4", "Oct 29", "Adaptation and integration: text, Cataract-101, point clouds"),
    ("Week 5", "Nov 5", "Downstream tasks and ablations"),
    ("Week 6", "Nov 12", "Reproducibility and delivery"),
]

CATARACT_CITE = "Cataract-101: Schoeffmann et al., MMSys 2018. Licence CC BY-NC 4.0 (research use only)."


# ---------------------------------------------------------------- data


def load(n):
    """Rows of reports/tables/table_NN.csv as dicts (values stay strings)."""
    with open(TABLES / f"table_{n:02d}.csv", newline="") as f:
        return list(csv.DictReader(f))


def row(rows, **match):
    """The single row whose columns equal the given values."""
    hits = [r for r in rows if all(r[k] == v for k, v in match.items())]
    assert len(hits) == 1, (match, len(hits))
    return hits[0]


def f(v, d=3):
    """Format a CSV value the way reports/RESULTS.md does (3 decimals; '–' when empty)."""
    if v in ("", None, "–"):
        return "–"
    x = float(v)
    if "." not in v and "e" not in v.lower():
        return v                          # integer column (e.g. active codes)
    if abs(x) >= 1e5:
        return f"{x:.3g}"
    return f"{x:.{d}f}"


def n(v):
    """A count stored as float (e.g. '56.0') shown without decimals; refuses non-integers."""
    x = float(v)
    assert x == int(x), v
    return str(int(x))


def pct(v):
    """Utilisation fraction -> percent, 1 decimal (a unit change, not a new number)."""
    return f"{float(v) * 100:.1f} %"


# ---------------------------------------------------------------- deck


class Deck:
    def __init__(self, week, date, out):
        self.week, self.date, self.out = week, date, out
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.W, self.H = self.prs.slide_width, self.prs.slide_height
        self.n = 0

    def slide(self, title=None, source=None):
        self.n += 1
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if title:
            header(s, title, self.n, self.week, self.date, self.W)
        if source:
            text(s, Inches(0.5), Inches(7.05), Inches(11.5), Inches(0.35), "Source: " + source, size=10,
                 color=GREY)
        return s

    def title_slide(self, subtitle, long_date):
        s = self.slide()
        box(s, 0, 0, self.W, self.H, fill=NAVY)
        box(s, Inches(0.8), Inches(2.0), Inches(0.12), Inches(2.6), fill=TEAL)
        text(s, Inches(1.2), Inches(1.9), Inches(11), Inches(0.5),
             f"UNDERGRADUATE THESIS · WEEKLY PROGRESS REPORT {self.week}", size=16, color=MINT, bold=True)
        text(s, Inches(1.2), Inches(2.4), Inches(11.3), Inches(1.8),
             "Design of a Medical Multimodal-Specific Tokenizer and a Unified Data Protocol", size=40,
             color=WHITE, bold=True)
        text(s, Inches(1.2), Inches(4.6), Inches(11), Inches(0.6), subtitle, size=20, color=PALE)
        text(s, Inches(1.2), Inches(5.7), Inches(11), Inches(1.0),
             [("REDOANUZZAMAN", {"bold": True, "size": 20, "color": WHITE}),
              (f"Sichuan University  ·  {long_date}", {"size": 16, "color": PALE})])
        return s

    def save(self):
        self.prs.save(self.out)
        print(self.out, f"({self.n} slides)")


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, lines, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, bullet=False, spacing=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        # a line may be (text, overrides) to vary size / colour / boldness
        t, o = (ln, {}) if isinstance(ln, str) else ln
        r = p.add_run()
        r.text = ("•  " if bullet and not o.get("nobullet") else "") + t
        r.font.name = o.get("font", font)
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.color.rgb = o.get("color", color)
    return tb


def header(slide, title, n, week, date, W):
    box(slide, 0, 0, W, Inches(1.0), fill=NAVY)
    text(slide, Inches(0.5), Inches(0.12), Inches(10), Inches(0.8), title, size=28, color=WHITE,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(10.3), Inches(0.12), Inches(2.6), Inches(0.8), f"Week {week}  ·  {date}, 2026",
         size=14, color=PALE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(12.2), Inches(7.0), Inches(0.9), Inches(0.4), str(n), size=12, color=GREY,
         align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=TEAL, size=15):
    box(slide, x, y, w, h, fill=LIGHT)
    box(slide, x, y, Inches(0.08), h, fill=accent)
    text(slide, x + Inches(0.25), y + Inches(0.12), w - Inches(0.35), Inches(0.5), title, size=18,
         bold=True, color=accent)
    text(slide, x + Inches(0.25), y + Inches(0.62), w - Inches(0.35), h - Inches(0.7), body,
         size=size, bullet=True, spacing=4)


def table(slide, x, y, w, rows, col_w, size=14, row_h=Inches(0.45), bold_rows=()):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows)).table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i in range(len(rows)):
        t.rows[i].height = row_h
    for i, r_ in enumerate(rows):
        for j, val in enumerate(r_):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            r = tf.paragraphs[0].add_run()
            r.text = val
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold = i == 0 or i in bold_rows
            r.font.color.rgb = WHITE if i == 0 else DARK
    return t


def figure(slide, name, x, y, max_w, max_h, align="center"):
    """Place reports/figures/<name>.png as large as fits in the box, keeping its aspect ratio."""
    path = FIGS / f"{name}.png"
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    dx = (max_w - w) // 2 if align == "center" else 0
    slide.shapes.add_picture(str(path), x + dx, y + (max_h - h) // 2, w, h)
    return w, h


def banner(slide, y, msg, accent=AMBER, x=Inches(0.5), w=Inches(12.3), h=Inches(0.55), size=14):
    """A thin warning strip for the conditions that must stay visible (WEEKLY_PLAN.md)."""
    box(slide, x, y, w, h, fill=RGBColor(0xFE, 0xF3, 0xC7))
    box(slide, x, y, Inches(0.08), h, fill=accent)
    text(slide, x + Inches(0.2), y, w - Inches(0.3), h, msg, size=size, color=DARK,
         anchor=MSO_ANCHOR.MIDDLE)


def roadmap(slide, current, y0=Inches(1.3), step=0.62):
    """Compact six-week strip with done / current / upcoming weeks."""
    for i, (wk, date, d) in enumerate(PLAN):
        y = y0 + Inches(i * step)
        k = i + 1
        fill = TEAL if k == current else (NAVY if k < current else LIGHT)
        box(slide, Inches(0.5), y, Inches(1.5), Inches(step - 0.1), fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(slide, Inches(0.5), y, Inches(1.5), Inches(step - 0.1), wk, size=14, bold=True,
             color=WHITE if k <= current else NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        label = d + ("   ✓ done" if k < current else ("   ◀ now" if k == current else ""))
        text(slide, Inches(2.1), y, Inches(4.55), Inches(step - 0.1), label, size=13, bold=k == current,
             color=DARK if k <= current else GREY, anchor=MSO_ANCHOR.MIDDLE)
